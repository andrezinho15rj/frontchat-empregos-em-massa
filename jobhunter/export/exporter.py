"""Exportacao de dados para CSV e Excel (.xlsx)."""
from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from .. import app_paths
from ..db.database import database
from ..logging_setup import get_logger

log = get_logger("export")

EXPORT_COLUMNS = [
    ("empresa", "Empresa"),
    ("cargo", "Cargo"),
    ("cidade", "Cidade"),
    ("estado", "Estado"),
    ("email", "E-mail"),
    ("site", "Site da empresa"),
    ("url_vaga", "URL da vaga"),
    ("fonte_contato", "Fonte do contato"),
    ("data_publicacao", "Data de publicação"),
    ("ultima_verificacao", "Data da última verificação"),
    ("status_candidatura", "Status da candidatura"),
]


def _query(status_filter: str = "", area_filter: str = "", search: str = "",
           only_with_email: bool = False) -> list[dict]:
    """Monta linhas unindo vaga + empresa + contato + candidatura."""
    clauses, params = [], []
    if only_with_email:
        clauses.append("ct.email IS NOT NULL")
    if status_filter:
        clauses.append("a.status = ?")
        params.append(status_filter)
    if area_filter:
        clauses.append("j.area = ?")
        params.append(area_filter)
    if search:
        like = f"%{search}%"
        clauses.append("(c.name LIKE ? OR j.title LIKE ? OR ct.email LIKE ?)")
        params += [like, like, like]
    where = ("WHERE " + " AND ".join(clauses)) if clauses else ""

    sql = f"""
        SELECT
            c.name AS empresa,
            j.title AS cargo,
            COALESCE(NULLIF(j.city,''), c.location) AS cidade,
            j.state AS estado,
            ct.email AS email,
            c.website AS site,
            j.url AS url_vaga,
            ct.source_url AS fonte_contato,
            j.published_at AS data_publicacao,
            ct.checked_at AS ultima_verificacao,
            a.status AS status_candidatura
        FROM jobs j
        LEFT JOIN companies c ON c.id = j.company_id
        LEFT JOIN contacts ct ON ct.job_id = j.id OR ct.company_id = j.company_id
        LEFT JOIN applications a ON a.job_id = j.id AND a.to_email = ct.email
        {where}
        ORDER BY j.id DESC
        LIMIT 20000
    """
    rows = database.query(sql, tuple(params))
    result = []
    seen = set()
    for r in rows:
        key = (r["empresa"], r["cargo"], r["email"])
        if key in seen:
            continue
        seen.add(key)
        result.append({k: (r[k] if r[k] is not None else "") for k, _ in EXPORT_COLUMNS})
    return result


class Exporter:
    def preview(self, **filters) -> list[dict]:
        return _query(**filters)

    def to_csv(self, rows: list[dict], path: str | Path | None = None) -> Path:
        path = Path(path) if path else app_paths.exports_dir() / \
            f"jobhunter_{datetime.now():%Y%m%d_%H%M%S}.csv"
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.writer(fh, delimiter=";")
            writer.writerow([label for _, label in EXPORT_COLUMNS])
            for row in rows:
                writer.writerow([row.get(key, "") for key, _ in EXPORT_COLUMNS])
        log.info("Exportado CSV: %s (%d linhas)", path, len(rows))
        return path

    def to_xlsx(self, rows: list[dict], path: str | Path | None = None) -> Path:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        path = Path(path) if path else app_paths.exports_dir() / \
            f"jobhunter_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Candidaturas"
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="1F4E78")
        ws.append([label for _, label in EXPORT_COLUMNS])
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
        for row in rows:
            ws.append([row.get(key, "") for key, _ in EXPORT_COLUMNS])
        for col_idx, (_, label) in enumerate(EXPORT_COLUMNS, start=1):
            width = max(len(label) + 2, 14)
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width
        wb.save(str(path))
        log.info("Exportado XLSX: %s (%d linhas)", path, len(rows))
        return path


exporter = Exporter()
