"""Testes de importação de currículo, templates e preparação de campanha."""
from __future__ import annotations

import unittest

from .base import DBTestCase


def _make_docx(path):
    import docx
    d = docx.Document()
    d.add_paragraph("João da Silva")
    d.add_paragraph("Analista Administrativo com experiência em rotinas de RH.")
    d.add_paragraph("Competências: Excel, atendimento, organização de documentos.")
    d.save(str(path))


class ResumeTests(DBTestCase):
    def test_import_docx(self):
        import tempfile
        from pathlib import Path
        from jobhunter.resume import import_resume, candidate_name

        p = Path(tempfile.mkdtemp()) / "cv.docx"
        _make_docx(p)
        rid, text, ftype = import_resume(p, name="João da Silva")
        self.assertGreater(len(text), 20)
        self.assertEqual(ftype, "docx")
        self.assertIn("João da Silva", candidate_name(text))


class TemplateTests(unittest.TestCase):
    def test_render_variables(self):
        from jobhunter.campaigns.templates import render, build_values
        values = build_values("Empresa XYZ", "Analista Administrativo", "João",
                              city="São Paulo")
        out = render("Candidatura para {cargo} — {candidato} ({empresa})", values)
        self.assertEqual(out, "Candidatura para Analista Administrativo — João (Empresa XYZ)")

    def test_missing_variable_is_empty(self):
        from jobhunter.campaigns.templates import render
        self.assertEqual(render("Olá {inexistente}!", {}), "Olá !")


class CampaignPrepareTests(DBTestCase):
    def test_prepare_creates_applications_without_sending(self):
        import tempfile
        from pathlib import Path
        from jobhunter.resume import import_resume
        from jobhunter.db.repositories import (
            application_repo, campaign_repo, company_repo, contact_repo, job_repo,
        )
        from jobhunter.campaigns import campaign_engine, default_templates

        p = Path(tempfile.mkdtemp()) / "cv.docx"
        _make_docx(p)
        rid, _, _ = import_resume(p, name="João da Silva")

        cid = company_repo.get_or_create("Empresa XPTO", domain="xpto.com.br")
        job_id = job_repo.upsert({"company_id": cid, "title": "Analista Administrativo",
                                  "url": "https://xpto.com.br/v/1", "source": "T",
                                  "area": "Administrativo"})
        contact_id = contact_repo.add(cid, email="rh@xpto.com.br", domain="xpto.com.br",
                                      source_url="https://xpto.com.br/contatos", job_id=job_id)
        s, b = default_templates("pt")
        camp_id = campaign_repo.create("Teste", rid, s, b, "pt")

        report = campaign_engine.prepare(camp_id, [contact_id], use_ai=False)
        self.assertEqual(report.prepared, 1)
        apps = application_repo.list(campaign_id=camp_id)
        self.assertEqual(apps[0]["status"], "Preparada")
        self.assertIn("Analista Administrativo", apps[0]["subject"])
        self.assertIn("Empresa XPTO", apps[0]["body"])
        self.assertIn("João", apps[0]["body"])

    def test_no_duplicate_application_for_same_job(self):
        import tempfile
        from pathlib import Path
        from jobhunter.resume import import_resume
        from jobhunter.db.repositories import (
            application_repo, campaign_repo, company_repo, contact_repo, job_repo,
        )
        from jobhunter.campaigns import campaign_engine, default_templates

        p = Path(tempfile.mkdtemp()) / "cv.docx"
        _make_docx(p)
        rid, _, _ = import_resume(p, name="Maria")
        cid = company_repo.get_or_create("Emp", domain="emp.com")
        job_id = job_repo.upsert({"company_id": cid, "title": "Assistente",
                                  "url": "https://emp.com/v/9", "source": "T"})
        contact_id = contact_repo.add(cid, email="rh@emp.com", source_url="https://emp.com", job_id=job_id)
        s, b = default_templates("pt")
        camp = campaign_repo.create("T2", rid, s, b, "pt")

        # simula envio anterior
        aid = application_repo.add(campaign_id=camp, company_id=cid, job_id=job_id,
                                   contact_id=contact_id, to_email="rh@emp.com",
                                   status="Enviada")
        report = campaign_engine.prepare(camp, [contact_id], use_ai=False)
        self.assertEqual(report.prepared, 0)
        self.assertGreaterEqual(report.skipped, 1)


if __name__ == "__main__":
    unittest.main()
